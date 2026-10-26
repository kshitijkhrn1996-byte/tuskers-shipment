"""Read the judges' Bookings (Sheet 1) and Items (Sheet 2) data.

Accepts one workbook with both sheets, or several files (xlsx or csv).
Each sheet or file is recognised by its header row, not its name, so
renamed sheets and the template's other sheets (rate card, rulebook,
guides) don't matter. Headers are matched loosely (case, spacing, units
in brackets).
"""
import csv
import io
import re
from typing import Any, Iterable, Optional

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from .schemas import Booking

BOOKING_COLS = {
    "bookingid": "booking_id", "sellername": "seller_name", "origincity": "origin_city",
    "origincountry": "origin_country", "destinationcity": "destination_city",
    "destinationcountry": "destination_country", "destinationtype": "destination_type",
    "readytoshipdate": "ready_to_ship_date", "sellerpriority": "seller_priority",
    "pallets": "pallets", "cartons": "cartons", "totaldeclaredvalue": "total_declared_value",
    "currency": "currency", "totalgrossweight": "total_gross_weight_kg",
    "deliveryappointmentbooked": "delivery_appointment_booked",
}
ITEM_COLS = {
    "bookingid": "booking_id", "sku": "sku", "description": "description", "hscode": "hs_code",
    "quantity": "quantity", "unitprice": "unit_price", "unitspercarton": "units_per_carton",
    "cartons": "cartons", "unitweight": "unit_weight_kg", "countryoforigin": "country_of_origin",
}

# Row 3 of the template is a grey example. Judges enter data from row 4, so
# skip that row only if it is left in exactly as shipped.
EXAMPLE_ROWS = {
    ("BK-001", "Northfield Goods", "Manchester", "United Kingdom", "Frankfurt", "Germany", "Warehouse",
     "05-11-2026", "Cheapest", "3", "90", "5400", "GBP", "620", "Yes"),
    ("BK-001", "MUG-01", "Ceramic coffee mug, 350 ml", "691200", "1080", "5", "12", "90", "0.5", "United Kingdom"),
}


class RowError(BaseModel):
    booking_id: str
    errors: list[str]


def _norm(h: Any) -> str:
    h = re.sub(r"\(.*?\)", "", str(h or "")).lower()
    return re.sub(r"[^a-z]", "", h)


def _cell(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v is not None else ""


def _canon(v) -> str:
    """Cell text with numbers normalised ("5.00" == 5.0 == "5") for comparisons."""
    s = _cell(v)
    try:
        f = float(s)
    except ValueError:
        return s
    return str(int(f)) if f.is_integer() else f"{f:g}"


def _is_example(row) -> bool:
    return tuple(_canon(v) for v in row if _cell(v)) in EXAMPLE_ROWS


def _classify(row) -> Optional[tuple[str, dict[int, str]]]:
    """Is this a Bookings or Items header row? Returns (kind, column index -> field)."""
    b = {i: BOOKING_COLS[_norm(c)] for i, c in enumerate(row) if _norm(c) in BOOKING_COLS}
    if "booking_id" in b.values() and {"seller_name", "destination_city", "pallets"} & set(b.values()):
        return "bookings", b
    it = {i: ITEM_COLS[_norm(c)] for i, c in enumerate(row) if _norm(c) in ITEM_COLS}
    if "booking_id" in it.values() and {"sku", "units_per_carton", "hs_code"} & set(it.values()):
        return "items", it
    return None


def _scan(rows: Iterable) -> tuple[Optional[str], list[dict]]:
    kind, idx, out = None, None, []
    for row in rows:
        row = tuple(row)
        if idx is None:
            found = _classify(row)
            if found:
                kind, idx = found
            continue
        if all(not _cell(v) for v in row) or _is_example(row):
            continue
        out.append({field: (row[i] if row[i] != "" else None) for i, field in idx.items() if i < len(row)})
    return kind, out


def _tables(name: str, data: bytes):
    if name.lower().endswith(".csv"):
        yield csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
        return
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        yield ws.iter_rows(values_only=True)


def _fmt(e: ValidationError) -> list[str]:
    msgs = []
    for err in e.errors():
        loc = ".".join(str(p) for p in err["loc"])
        msgs.append(f"{loc}: {err['msg']}" if loc else err["msg"])
    return msgs


def parse_files(files: list[tuple[str, bytes]]) -> tuple[list[Booking], list[RowError]]:
    bookings_raw, items_raw, kinds = [], [], set()
    for name, data in files:
        for rows in _tables(name, data):
            kind, records = _scan(rows)
            kinds.add(kind)
            if kind == "bookings":
                bookings_raw += records
            elif kind == "items":
                items_raw += records
    if "bookings" not in kinds:
        raise ValueError("Couldn't find a Bookings table (a header row with Booking ID, Seller name, "
                         "Destination city, Pallets …). Upload the Bookings and Items sheets.")

    by_booking: dict[str, list[dict]] = {}
    for it in items_raw:
        by_booking.setdefault(_cell(it.get("booking_id")), []).append(it)

    bookings, errors, seen = [], [], set()
    for raw in bookings_raw:
        bid = _cell(raw.get("booking_id"))
        if bid in seen:
            errors.append(RowError(booking_id=bid, errors=["Duplicate Booking ID — each booking needs its own ID; "
                                                           "this repeated row was skipped."]))
            continue
        seen.add(bid)
        raw["items"] = by_booking.get(bid, [])
        try:
            b = Booking.model_validate(raw)
        except ValidationError as e:
            errors.append(RowError(booking_id=bid or "?", errors=_fmt(e)))
            continue
        if not b.items:
            errors.append(RowError(booking_id=bid, errors=["No products found for this booking on the Items sheet."]))
            continue
        bookings.append(b)
    for bid in by_booking:
        if bid and bid not in seen:
            errors.append(RowError(booking_id=bid, errors=[f"{len(by_booking[bid])} product row(s) on the Items sheet "
                                                           "use this Booking ID, but there's no such booking on the "
                                                           "Bookings sheet."]))
    return bookings, errors


def parse_workbook(data: bytes) -> tuple[list[Booking], list[RowError]]:
    return parse_files([("upload.xlsx", data)])
